from io import BytesIO
from datetime import date
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x0A, 0x0F, 0x1E)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GRAY = RGBColor(0x5A, 0x61, 0x72)
LOGO_PATH = "/app/backend/assets/letterhead_logo.png"

GREEK_MONTHS = ["", "Ιανουαρίου", "Φεβρουαρίου", "Μαρτίου", "Απριλίου", "Μαΐου", "Ιουνίου",
                "Ιουλίου", "Αυγούστου", "Σεπτεμβρίου", "Οκτωβρίου", "Νοεμβρίου", "Δεκεμβρίου"]


def _gold_rule(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C9A227")
    pBdr.append(bottom)
    pPr.append(pBdr)


def build_letter_docx(settings: dict, recipient: str, subject: str, body: str) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(1.5), Cm(1.5)
        s.left_margin, s.right_margin = Cm(2.3), Cm(2.3)

    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")

    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.add_run().add_picture(LOGO_PATH, width=Cm(5.4))

    rule = doc.add_paragraph()
    _gold_rule(rule)

    today = date.today()
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    dr = date_p.add_run(f"Αθήνα, {today.day} {GREEK_MONTHS[today.month]} {today.year}")
    dr.font.size = Pt(10)
    dr.font.color.rgb = GRAY

    if recipient:
        rec_p = doc.add_paragraph()
        rr = rec_p.add_run(f"Προς: {recipient}")
        rr.font.size = Pt(11)
        rr.bold = True
        rr.font.color.rgb = NAVY

    if subject:
        sub_p = doc.add_paragraph()
        sr = sub_p.add_run(f"Θέμα: {subject}")
        sr.bold = True
        sr.font.size = Pt(12)
        sr.font.color.rgb = NAVY

    doc.add_paragraph()
    for para in [p.strip() for p in body.replace("\r", "").split("\n") if p.strip()]:
        bp = doc.add_paragraph(para)
        bp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        bp.paragraph_format.space_after = Pt(8)

    doc.add_paragraph()
    close_p = doc.add_paragraph("Με εκτίμηση,")
    close_p.paragraph_format.space_after = Pt(2)
    comp_p = doc.add_paragraph()
    cr = comp_p.add_run(settings.get("company_name") or "Nomos One")
    cr.bold = True
    cr.font.color.rgb = NAVY

    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parts = [settings.get("address"), f"Τηλ. {settings['phone']}" if settings.get("phone") else None, settings.get("email")]
    fr = fp.add_run("  |  ".join([p for p in parts if p]))
    fr.font.size = Pt(8.5)
    fr.font.color.rgb = GRAY
    _gold_rule(fp)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
